"""文档类提取器：PDF / DOCX / PPTX / XLSX / RTF / 纯文本。

所有第三方库惰性导入；缺失时抛 :class:`ExtractError`，由 dispatch 转成
``ExtractionResult(ok=False, error=...)``。

每个提取器的签名统一为::

    handler(path: str, *, ocr_lang: str = "ch", whisper_model: str = "small", **kw)
        -> (text: str, meta: dict, method: str)
"""

from __future__ import annotations

import codecs
import json
import re
from pathlib import Path

from ._common import ExtractError, looks_binary, require
from .models import MAX_TEXT_BYTES

# --------------------------------------------------------------------------- #
# 编码探测
# --------------------------------------------------------------------------- #
#: 依次尝试的编码（宽到窄，最后是几乎不会失败的兜底）
_ENCODING_CHAIN: tuple[str, ...] = (
    "utf-8-sig",
    "utf-8",
    "gb18030",
    "big5",
    "cp932",
    "koi8-r",
    "cp1252",
    "latin-1",
)


def _decode_bytes(data: bytes) -> tuple[str, str]:
    """按编码链解码，返回 ``(text, encoding)``。"""
    if data.startswith(codecs.BOM_UTF8):
        return data.decode("utf-8-sig", errors="replace"), "utf-8-sig"
    if data.startswith(codecs.BOM_UTF16_LE) or data.startswith(codecs.BOM_UTF16_BE):
        return data.decode("utf-16", errors="replace"), "utf-16"
    for enc in _ENCODING_CHAIN:
        try:
            return data.decode(enc), enc
        except (UnicodeDecodeError, LookupError):
            continue
    return data.decode("latin-1", errors="replace"), "latin-1"


# --------------------------------------------------------------------------- #
# 纯文本
# --------------------------------------------------------------------------- #
def extract_text(path: str, **_kw) -> tuple[str, dict, str]:
    """直读文本文件（编码探测）。二进制内容判定为失败。"""
    raw = Path(path).read_bytes()
    if looks_binary(raw):
        raise ExtractError(
            f"`{Path(path).name}` 看起来是二进制文件（含 NUL/大量控制字符），不是文本文件，已跳过。"
        )
    text, encoding = _decode_bytes(raw)
    newline = "\r\n" if b"\r\n" in raw[:4096] else "\n"
    meta = {
        "encoding": encoding,
        "bytes": len(raw),
        "lines": text.count("\n") + (1 if text and not text.endswith("\n") else 0),
        "newline": "CRLF" if newline == "\r\n" else "LF",
    }
    if path.lower().endswith(".json"):
        try:
            json.loads(text)
            meta["json_valid"] = True
        except Exception as exc:
            meta["json_valid"] = False
            meta["json_error"] = f"{type(exc).__name__}: {exc}"
    return text, meta, "passthrough"


# --------------------------------------------------------------------------- #
# PDF
# --------------------------------------------------------------------------- #
def extract_pdf(path: str, **_kw) -> tuple[str, dict, str]:
    pypdf = require("pypdf", "pypdf", "pip install pypdf", "PDF 文本抽取")

    try:
        reader = pypdf.PdfReader(path)
    except Exception as exc:
        raise ExtractError(f"PDF 无法打开（可能损坏或加密）：{type(exc).__name__}: {exc}") from exc

    if getattr(reader, "is_encrypted", False):
        try:
            if reader.decrypt("") == 0:
                raise ExtractError("PDF 已加密且空密码无法解密，请先解密后再提取。")
        except ExtractError:
            raise
        except Exception as exc:
            raise ExtractError(f"PDF 已加密，解密失败：{type(exc).__name__}: {exc}") from exc

    parts: list[str] = []
    page_stats: list[int] = []
    page_errors: list[str] = []
    for idx, page in enumerate(reader.pages, start=1):
        try:
            page_text = page.extract_text() or ""
        except Exception as exc:  # 单页失败不影响整体
            page_text = ""
            page_errors.append(f"page {idx}: {type(exc).__name__}: {exc}")
        page_stats.append(len(page_text.strip()))
        parts.append(f"[page {idx}]\n{page_text.strip()}")
    text = "\n\n".join(parts).strip()

    meta: dict = {
        "pages": len(reader.pages),
        "page_char_counts": page_stats,
        "has_text_layer": any(c > 0 for c in page_stats),
    }
    try:
        info = reader.metadata or {}
        meta["pdf_metadata"] = {str(k).lstrip("/"): str(v) for k, v in dict(info).items() if v}
    except Exception:
        meta["pdf_metadata"] = {}
    if page_errors:
        meta["page_errors"] = page_errors
    if not meta["has_text_layer"]:
        meta["note"] = (
            "PDF 无文本层（疑似扫描件）。本层不做页面渲染 OCR；"
            "如需内容，请先用外部工具把页面渲染成图片，再走 images 提取器做本地 OCR。"
        )
    return text, meta, "pypdf"


# --------------------------------------------------------------------------- #
# DOCX
# --------------------------------------------------------------------------- #
def _iter_docx_blocks(document):
    """按文档顺序产出 ('p', paragraph) / ('t', table)。

    python-docx >= 1.1 提供 ``iter_inner_content``；旧版退回「先段落后表格」。
    """
    try:
        from docx.table import Table
        from docx.text.paragraph import Paragraph

        for block in document.iter_inner_content():
            if isinstance(block, Paragraph):
                yield "p", block
            elif isinstance(block, Table):
                yield "t", block
        return
    except Exception:
        pass
    for paragraph in document.paragraphs:
        yield "p", paragraph
    for table in document.tables:
        yield "t", table


def extract_docx(path: str, **_kw) -> tuple[str, dict, str]:
    docx = require("docx", "python-docx", "pip install python-docx", "DOCX 文本抽取")

    try:
        document = docx.Document(path)
    except Exception as exc:
        raise ExtractError(f"DOCX 无法打开：{type(exc).__name__}: {exc}") from exc

    parts: list[str] = []
    paragraphs = 0
    tables = 0
    table_rows = 0
    for kind, block in _iter_docx_blocks(document):
        if kind == "p":
            value = (block.text or "").strip()
            if not value:
                continue
            style = ""
            try:
                style = block.style.name or ""
            except Exception:
                style = ""
            if style.lower().startswith("heading"):
                value = f"{'#' * min(6, max(1, _heading_level(style)))} {value}"
            parts.append(value)
            paragraphs += 1
        else:
            tables += 1
            rows: list[str] = []
            for row in block.rows:
                cells = [(cell.text or "").strip().replace("\n", " ") for cell in row.cells]
                if any(cells):
                    rows.append(" | ".join(cells))
            if rows:
                table_rows += len(rows)
                parts.append(f"[表格 {tables}]\n" + "\n".join(rows))

    text = "\n".join(parts)
    meta = {
        "paragraphs": paragraphs,
        "tables": tables,
        "table_rows": table_rows,
        "sections": len(getattr(document, "sections", []) or []),
    }
    return text, meta, "python-docx"


def _heading_level(style_name: str) -> int:
    match = re.search(r"(\d+)", style_name or "")
    return int(match.group(1)) if match else 1


# --------------------------------------------------------------------------- #
# PPTX（必须逐页，含表格与备注）
# --------------------------------------------------------------------------- #
def _shape_texts(shape) -> list[str]:
    out: list[str] = []
    has_table = bool(getattr(shape, "has_table", False))
    if has_table:
        try:
            for row in shape.table.rows:
                cells = [(cell.text or "").strip().replace("\n", " ") for cell in row.cells]
                if any(cells):
                    out.append("  ".join(cells))
        except Exception:
            pass
    if getattr(shape, "has_text_frame", False):
        for paragraph in shape.text_frame.paragraphs:
            value = (paragraph.text or "").strip()
            if value:
                out.append(value)
    if not has_table and not getattr(shape, "has_text_frame", False) and getattr(shape, "has_chart", False):
        try:
            chart = shape.chart
            out.append(f"[图表 {chart.chart_type}]")
        except Exception:
            pass
    return out


def extract_pptx(path: str, **_kw) -> tuple[str, dict, str]:
    pptx = require("pptx", "python-pptx", "pip install python-pptx", "PPTX 逐页抽取")

    try:
        presentation = pptx.Presentation(path)
    except Exception as exc:
        raise ExtractError(f"PPTX 无法打开：{type(exc).__name__}: {exc}") from exc

    slides: list[str] = []
    slide_meta: list[dict] = []
    notes_count = 0
    for idx, slide in enumerate(presentation.slides, start=1):
        lines: list[str] = []
        for shape in slide.shapes:
            lines.extend(_shape_texts(shape))
        note_text = ""
        try:
            if slide.has_notes_slide:
                frame = slide.notes_slide.notes_text_frame
                note_text = (frame.text or "").strip() if frame is not None else ""
        except Exception:
            note_text = ""
        body = "\n".join(lines)
        chunk = f"[slide {idx}]\n{body}"
        if note_text:
            notes_count += 1
            chunk += f"\n[备注]\n{note_text}"
        slides.append(chunk.strip())
        slide_meta.append(
            {
                "slide": idx,
                "shapes": len(slide.shapes),
                "chars": len(body),
                "has_notes": bool(note_text),
            }
        )

    text = "\n\n".join(slides).strip()
    meta = {
        "slides": len(presentation.slides),
        "slide_stats": slide_meta,
        "notes_slides": notes_count,
    }
    return text, meta, "python-pptx"


# --------------------------------------------------------------------------- #
# XLSX / XLSM
# --------------------------------------------------------------------------- #
def extract_xlsx(path: str, **_kw) -> tuple[str, dict, str]:
    openpyxl = require("openpyxl", "openpyxl", "pip install openpyxl", "Excel 表格抽取")

    try:
        workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    except Exception as exc:
        raise ExtractError(f"Excel 无法打开：{type(exc).__name__}: {exc}") from exc

    parts: list[str] = []
    sheets_meta: list[dict] = []
    try:
        for sheet in workbook.worksheets:
            rows: list[str] = []
            for row in sheet.iter_rows(values_only=True):
                cells = ["" if v is None else str(v).replace("\n", " ") for v in row]
                while cells and cells[-1] == "":
                    cells.pop()
                if any(c.strip() for c in cells):
                    rows.append(" | ".join(cells))
            parts.append(f"[sheet {sheet.title}]\n" + "\n".join(rows))
            sheets_meta.append({"name": sheet.title, "rows": len(rows)})
    finally:
        try:
            workbook.close()
        except Exception:
            pass

    text = "\n\n".join(parts).strip()
    meta = {"sheets": sheets_meta, "sheet_count": len(sheets_meta)}
    if not text:
        meta["note"] = "工作簿为空，或所有单元格值均未被缓存（data_only=True 只读取公式结果）。"
    return text, meta, "openpyxl"


# --------------------------------------------------------------------------- #
# RTF（尽力而为）
# --------------------------------------------------------------------------- #
#: 需要整体丢弃的 RTF 目标组（字体表、颜色表、图片、字段等）
_RTF_DROP_DESTS = frozenset(
    {
        "fonttbl",
        "colortbl",
        "stylesheet",
        "info",
        "pict",
        "object",
        "themedata",
        "latentstyles",
        "datastore",
        "generator",
        "listtable",
        "listoverridetable",
        "rsidtbl",
        "filetbl",
        "header",
        "footer",
        "footnote",
        "annotation",
        "field",
        "bkmkstart",
        "bkmkend",
    }
)
_RTF_UNICODE_RE = re.compile(r"\\u(-?\d+)\s?\??")
_RTF_HEX_RE = re.compile(r"\\'([0-9a-fA-F]{2})")
_RTF_CONTROL_RE = re.compile(r"\\[a-zA-Z]+-?\d*\s?")
_RTF_SYMBOL_RE = re.compile(r"\\[^a-zA-Z]")
_RTF_DEST_HEAD_RE = re.compile(r"\\([a-zA-Z]+)")


def _strip_rtf_dest_groups(data: str) -> str:
    """按花括号配对整体删除目标组（``{\\fonttbl ...}``、``{\\*\\generator ...}`` 等）。

    手写扫描器是必要的：``fonttbl``/``pict`` 内部还有嵌套花括号，正则匹配不了。
    """
    out: list[str] = []
    index = 0
    length = len(data)
    while index < length:
        char = data[index]
        if char != "{":
            out.append(char)
            index += 1
            continue

        cursor = index + 1
        while cursor < length and data[cursor] in " \t\r\n":
            cursor += 1

        drop = data.startswith("\\*", cursor)
        if not drop:
            match = _RTF_DEST_HEAD_RE.match(data, cursor)
            if match and match.group(1).lower() in _RTF_DROP_DESTS:
                drop = True
        if not drop:
            out.append(char)
            index += 1
            continue

        depth = 0
        walker = index
        while walker < length:
            current = data[walker]
            if current == "\\":
                walker += 2
                continue
            if current == "{":
                depth += 1
            elif current == "}":
                depth -= 1
                if depth == 0:
                    walker += 1
                    break
            walker += 1
        out.append(" ")
        index = walker
    return "".join(out)


def extract_rtf(path: str, **_kw) -> tuple[str, dict, str]:
    """RTF 尽力提取：剥控制字、还原 ``\\uN`` 与 ``\\'hh``。"""
    raw = Path(path).read_bytes()
    if looks_binary(raw) and not raw.lstrip().startswith(b"{\\rtf"):
        raise ExtractError(f"`{Path(path).name}` 不是有效的 RTF（缺少 {{\\rtf 头）。")
    data = raw.decode("latin-1", errors="replace")

    data = _strip_rtf_dest_groups(data)
    data = _RTF_UNICODE_RE.sub(lambda m: chr(int(m.group(1)) % 65536), data)
    data = _RTF_HEX_RE.sub(lambda m: chr(int(m.group(1), 16)), data)
    data = data.replace("\\par", "\n").replace("\\line", "\n").replace("\\tab", "\t")
    data = data.replace("\\{", "\x01").replace("\\}", "\x02").replace("\\\\", "\x03")
    data = _RTF_CONTROL_RE.sub("", data)
    data = _RTF_SYMBOL_RE.sub("", data)
    data = data.replace("{", "").replace("}", "")
    data = data.replace("\x01", "{").replace("\x02", "}").replace("\x03", "\\")
    data = re.sub(r"[ \t]+\n", "\n", data)
    data = re.sub(r"\n{3,}", "\n\n", data)
    text = re.sub(r"[ \t]{2,}", " ", data.strip())

    meta = {"bytes": len(raw), "truncated_by_limit": len(text.encode("utf-8")) > MAX_TEXT_BYTES}
    if not text:
        meta["note"] = "RTF 未解析出可见文本（可能是纯图片/嵌入对象文档）。"
    return text, meta, "rtf-strip"
